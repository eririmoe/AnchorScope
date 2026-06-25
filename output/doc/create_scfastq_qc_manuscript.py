from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUTDIR = ROOT / "output" / "doc"
DOCX_PATH = OUTDIR / "scfastq_qc_bioinformatics_manuscript_draft.docx"
MD_PATH = OUTDIR / "scfastq_qc_bioinformatics_manuscript_draft.md"


TITLE = (
    "scfastq-qc: structure-aware quality control and filtering for "
    "single-cell long-read FASTQ libraries"
)


MANUSCRIPT = [
    ("title", TITLE),
    (
        "paragraph",
        "[Author names and affiliations to be added]\n"
        "Correspondence: [corresponding author email]\n"
        "Manuscript type: Software / methods article for Bioinformatics",
    ),
    ("heading1", "Abstract"),
    ("heading2", "Motivation"),
    (
        "paragraph",
        "Single-cell long-read sequencing libraries preserve molecular context that "
        "is difficult to inspect with generic FASTQ quality-control tools alone. "
        "In these assays, read usability depends not only on length and base-level "
        "quality, but also on whether expected molecular features such as terminal "
        "adapters, cell-barcode-associated sequence, primers, or poly(A) tracts are "
        "present in the correct order and orientation. Existing generic long-read "
        "QC reports summarize sequence yield and quality distributions, but they "
        "do not directly evaluate whether each read conforms to a user-defined "
        "single-cell library architecture.",
    ),
    ("heading2", "Results"),
    (
        "paragraph",
        "We introduce scfastq-qc, a Python command-line tool for structure-aware "
        "quality control and read filtering of single-cell long-read FASTQ files. "
        "The method accepts a JSON description of expected anchor motifs and their "
        "ordering, searches each read and its reverse complement, assigns a "
        "library-structure class, and combines this information with read length, "
        "Phred quality, terminal-offset, ambiguous-base, and invalid-base checks. "
        "It produces self-contained interactive HTML reports, machine-readable "
        "JSON summaries, optional CSV exports, and optional passed/failed FASTQ "
        "files. On two 10,000-read example datasets included with the repository, "
        "scfastq-qc recovered full expected structure in 56.26% and 54.78% of reads, "
        "identified low-quality, missing-anchor, internal-adapter, truncation, and "
        "concatemer-like failure modes, and highlighted a strong orientation "
        "difference between the two examples. These outputs illustrate how "
        "structure-aware QC can diagnose library construction and sequencing "
        "issues before downstream single-cell long-read analysis.",
    ),
    ("heading2", "Availability and implementation"),
    (
        "paragraph",
        "scfastq-qc is implemented in Python 3.10+ with an optional Rust accelerator "
        "for fixed-anchor search. The current source code is available in the "
        "accompanying repository; a public repository URL, software license, "
        "archival DOI, and package-distribution information should be added before "
        "submission.",
    ),
    ("heading1", "Introduction"),
    (
        "paragraph",
        "Quality control is a critical checkpoint in sequencing-data analysis "
        "because undetected technical artifacts can propagate into alignment, "
        "quantification, isoform discovery, cell assignment, and downstream "
        "biological interpretation. For single-cell RNA-seq and related assays, "
        "QC reporting is often most useful when it reflects the structure of the "
        "assay rather than only generic yield and quality metrics. The Bioinformatics "
        "software-article model exemplified by QCatch frames QC tools around this "
        "need: a tool should consolidate complex processing outputs into clear, "
        "reproducible reports tailored to the data type and workflow being assessed.",
    ),
    (
        "paragraph",
        "Single-cell long-read sequencing creates an additional QC problem. Long "
        "reads can contain full-length or near-full-length molecules, but their "
        "interpretability depends on recovering the molecular architecture of the "
        "library. A usable read may be expected to contain a 5-prime adapter, "
        "barcode-adjacent sequence, transcript body, and 3-prime poly(A)-related "
        "signal in a specific order. Conversely, a read can have acceptable base "
        "quality but still be uninformative if terminal anchors are missing, "
        "anchors occur internally, motifs appear in an unexpected order, or the "
        "best explanation of the read is a reverse-complement orientation. These "
        "failure modes are not sufficiently captured by aggregate length, N50, "
        "or per-base quality summaries.",
    ),
    (
        "paragraph",
        "Generic tools such as FastQC and long-read-oriented tools such as NanoPlot "
        "are valuable for broad sequencing diagnostics, including quality-score "
        "and read-length distributions. However, they are not designed to evaluate "
        "a user-defined molecular grammar for single-cell long-read libraries. "
        "Similarly, library-specific preprocessing tools may trim or orient reads "
        "for a particular protocol, but the resulting decisions are not always "
        "available as an interactive, sample-level QC report that exposes structural "
        "failure modes and filtering consequences.",
    ),
    (
        "paragraph",
        "To address this gap, we developed scfastq-qc, a lightweight and configurable "
        "command-line tool for structure-aware FASTQ QC. The central design is to "
        "separate assay knowledge from the program logic: users define named anchors, "
        "matching rules, expected order, and QC thresholds in a JSON configuration "
        "file, while the software applies a consistent streaming analysis to each "
        "read. The output is intended to help users answer practical questions: "
        "whether expected motifs are detectable, whether they occur in the correct "
        "order, whether a library shows truncation or internal-adapter signatures, "
        "whether reads are mixed between orientations, and which reads should be "
        "retained for downstream analysis.",
    ),
    ("heading1", "Methods"),
    ("heading2", "Datasets"),
    (
        "paragraph",
        "The repository includes two example FASTQ files, each containing 10,000 "
        "reads, together with a JSON configuration defining two anchors: a fixed "
        "5-prime adapter sequence, CGACATGGCTACGATCCGACTT, allowing one mismatch, "
        "and a regular-expression poly(A) anchor, A{8,}. The expected structure "
        "places the 5-prime adapter before the poly(A) anchor. Example thresholds "
        "set the minimum long-read length to 200 bp and the minimum read Q-score "
        "to 13. These examples are used here as a demonstration dataset for report "
        "content and method behavior, not as a benchmark of biological performance.",
    ),
    ("heading2", "Design and implementation"),
    (
        "paragraph",
        "scfastq-qc is distributed as a Python package exposing a command-line "
        "entry point, scfastq-qc. The command line supports single-file execution "
        "and batch execution. In single-file mode, the user provides a FASTQ or "
        "FASTQ.GZ file, a JSON configuration, an output directory, and optional "
        "flags for CSV export and passed/failed FASTQ output. In batch mode, the "
        "tool accepts a directory, a file list, or a samples list embedded in the "
        "configuration file, and can process multiple samples concurrently. The "
        "software streams FASTQ records and therefore does not require loading the "
        "entire input file into memory.",
    ),
    (
        "paragraph",
        "The configuration model contains four main components: sample metadata, "
        "anchor definitions, expected structure, and QC thresholds. Anchors can be "
        "fixed sequences or regular expressions. Fixed anchors support approximate "
        "matching through a bounded mismatch parameter, while regex anchors support "
        "motifs such as homopolymer or degenerate sequence patterns. Thresholds "
        "control read-level decisions, including minimum read length, minimum read "
        "Q-score, maximum allowed terminal-anchor offset, high-N fraction, and "
        "sample-level warn/fail cutoffs for reporting.",
    ),
    (
        "paragraph",
        "For speed, scfastq-qc includes an optional Rust shared library for fixed "
        "anchor search. When available, the Python layer loads this library through "
        "ctypes. If the shared library is absent or cannot be built, the program "
        "falls back to a pure-Python implementation. The fallback uses exact search "
        "for zero-mismatch anchors and a segment-seeding strategy for approximate "
        "fixed-anchor search: a motif is split into max_mismatches + 1 segments, "
        "candidate positions are seeded by exact segment matches, and candidate "
        "windows are accepted only if their bounded Hamming distance is within "
        "the configured mismatch limit. Regex anchors are evaluated with Python "
        "regular expressions.",
    ),
    ("heading2", "Read-level structure classification"),
    (
        "paragraph",
        "For each FASTQ record, scfastq-qc searches for all configured anchors in "
        "the observed sequence and in the reverse complement. Anchor hits are sorted "
        "by genomic position and summarized into an ordered list of detected anchor "
        "names. The forward and reverse-complement interpretations are then scored "
        "by order validity, structure-label priority, number of detected anchors, "
        "number of hits, and total mismatch burden. The orientation with the higher "
        "score is retained as the best structural explanation of the read.",
    ),
    (
        "paragraph",
        "The classifier assigns mutually interpretable structure labels. Reads with "
        "all expected anchors in the expected order are labeled full_structure. "
        "Reads without anchor hits are labeled no_anchor_detected. Partial reads are "
        "labeled according to missing terminal anchors where possible, while repeated "
        "or internal motifs are assigned labels such as duplicated_anchor, "
        "internal_5p_anchor, internal_3p_anchor, or concatemer_candidate. This label "
        "set is designed to expose failure modes that are actionable during library "
        "construction review or downstream read filtering.",
    ),
    ("heading2", "Read-level QC buckets and sample-level verdicts"),
    (
        "paragraph",
        "Structure classification is combined with conventional FASTQ checks to "
        "produce a primary read-level QC bucket. The decision logic checks empty "
        "reads, minimum length, read Q-score, high-N fraction, invalid bases, missing "
        "anchors, unexpected anchor order, internal adapter evidence, concatemer-like "
        "patterns, terminal truncation, and dense anchor hits. A priority ordering "
        "assigns each read to one primary bucket, with pass reserved for reads that "
        "satisfy both basic quality and structural criteria.",
    ),
    (
        "paragraph",
        "At the sample level, scfastq-qc summarizes total reads, total bases, mean "
        "and median read length, N50, minimum and maximum read length, median read "
        "Q-score, long high-quality ratio, reversed-read ratio, truncation ratios, "
        "high-N ratio, invalid-base ratio, anchor detection ratios, anchor mismatch "
        "burden, structure counts, orientation counts, and QC-bucket counts. "
        "Configured warn/fail thresholds are applied to key metrics, including "
        "long high-quality ratio, correct anchor-order ratio, no-anchor ratio, "
        "reversed-read ratio, and high-N read ratio, yielding an overall pass, warn, "
        "or fail status.",
    ),
    ("heading2", "Reports and exports"),
    (
        "paragraph",
        "The primary output is a self-contained HTML report. It includes run "
        "configuration, threshold-based QC verdicts, overview cards, read-length "
        "and Q-score distributions, anchor detection and mismatch summaries, anchor "
        "position tracks, structure-classification plots, QC-bucket plots, and a "
        "read-structure heatmap. The report is generated without external web "
        "dependencies, enabling local review and archiving. The tool also writes "
        "summary.json for programmatic reuse. When requested, it exports summary "
        "CSV files, read-level CSV files, anchor-hit CSV files, passed.fastq, "
        "failed.fastq, and optionally FASTQs grouped by non-full structure class.",
    ),
    ("heading1", "Results"),
    ("heading2", "Example run summary"),
    (
        "paragraph",
        "We applied scfastq-qc to the two example 10,000-read FASTQ files included "
        "with the repository. Both samples completed successfully in batch mode, "
        "and the batch report summarized per-sample total reads, long high-quality "
        "ratio, correct-anchor-order ratio, and overall QC status. The first example "
        "contained 8.65 Mb of sequence, with a median read length of 745 bp, N50 of "
        "926 bp, and median read Q-score of 15.58. The second example contained "
        "7.48 Mb of sequence, with a median read length of 685 bp, N50 of 769 bp, "
        "and median read Q-score of 15.07.",
    ),
    ("table", "Table 1. Example dataset summary", [
        ["Metric", "real_10k_1", "real_10k_2"],
        ["Total reads", "10,000", "10,000"],
        ["Total bases", "8,650,110", "7,475,020"],
        ["Median read length", "745 bp", "685 bp"],
        ["N50", "926 bp", "769 bp"],
        ["Median read Q-score", "15.58", "15.07"],
        ["Long high-quality ratio", "86.06%", "82.23%"],
        ["Full-structure reads", "56.26%", "54.78%"],
        ["Reversed-read ratio", "2.85%", "49.17%"],
        ["Overall verdict", "warn", "warn"],
    ]),
    (
        "paragraph",
        "The demonstration also shows why structure-aware reporting is informative. "
        "Both examples passed the configured long-read quality criterion, with "
        "long high-quality ratios above 80%. However, both were assigned an overall "
        "warn status because the fraction of reads with correct anchor order was "
        "below the configured warning threshold. This distinction would be difficult "
        "to infer from length and quality distributions alone.",
    ),
    ("heading2", "Anchor detection and structural failure modes"),
    (
        "paragraph",
        "The 5-prime adapter was detected in 73.24% and 73.80% of reads in the two "
        "examples, while the poly(A) anchor was detected in 95.53% and 95.98% of "
        "reads. The mean best-hit mismatch burden for the fixed 5-prime adapter was "
        "approximately 0.125 in both examples, consistent with near-exact recovery "
        "under the configured one-mismatch rule. The poly(A) anchor showed higher "
        "multi-hit ratios, 19.02% and 21.20%, reflecting the expected ambiguity of "
        "homopolymer-like sequence patterns and the need to interpret motif "
        "multiplicity alongside detection rate.",
    ),
    (
        "paragraph",
        "Structure labels separated read populations into interpretable classes. "
        "In real_10k_1, 5,626 reads were classified as full_structure, 2,110 as "
        "missing_5p_anchor, 1,817 as internal_3p_anchor, 354 as missing_3p_anchor, "
        "and 93 as no_anchor_detected. In real_10k_2, 5,478 reads were classified "
        "as full_structure, 2,049 as missing_5p_anchor, 1,928 as internal_3p_anchor, "
        "292 as missing_3p_anchor, 108 as concatemer_candidate, 36 as "
        "internal_5p_anchor, 4 as anchor_order_invalid, and 105 as "
        "no_anchor_detected. These labels give users a direct vocabulary for "
        "distinguishing absent anchors, internal adapter-like signatures, and "
        "concatemer-like patterns.",
    ),
    ("heading2", "Orientation-aware QC"),
    (
        "paragraph",
        "The two examples differed markedly in orientation composition. In real_10k_1, "
        "only 2.85% of reads were best explained by the reverse-complement orientation, "
        "and nearly all full-structure reads were forward. In real_10k_2, 49.17% of "
        "reads were best explained as reversed, including 2,753 of 5,478 full-structure "
        "reads. This sample-level difference was captured by the reversed-read "
        "verdict, which remained pass for the first example and warn for the second. "
        "Such orientation-aware summaries are useful when sequencing or preprocessing "
        "protocols can yield mixed strand representations.",
    ),
    ("heading2", "Filtering outputs"),
    (
        "paragraph",
        "When passed and failed FASTQ export was enabled, scfastq-qc wrote 5,095 "
        "passed reads and 4,905 failed reads for real_10k_1, and 4,573 passed reads "
        "and 5,427 failed reads for real_10k_2. The dominant non-pass buckets were "
        "internal_adapter, missing_5p, and low_read_q in the first example, and "
        "low_read_q, internal_adapter, and missing_5p in the second. Because the "
        "program also writes machine-readable summaries, users can tune thresholds "
        "or inspect individual read-level calls before committing to downstream "
        "filtering.",
    ),
    ("heading1", "Discussion"),
    (
        "paragraph",
        "scfastq-qc addresses a practical gap in single-cell long-read workflows by "
        "making library structure a first-class QC object. Instead of treating FASTQ "
        "QC as only a question of length and base quality, it asks whether reads "
        "carry the expected molecular features in the expected orientation and "
        "order. This design is especially useful for assays in which successful "
        "downstream interpretation depends on recovering terminal adapters, "
        "barcode-proximal segments, transcript bodies, and poly(A)-associated "
        "sequence features from noisy long reads.",
    ),
    (
        "paragraph",
        "The method is intentionally configurable. Users can adapt anchor definitions "
        "and expected order to different library designs without changing the source "
        "code. Fixed anchors support approximate matching, which is important for "
        "long-read data where substitutions can occur, while regex anchors allow "
        "simple representation of motifs such as poly(A) tracts. The combination of "
        "orientation-aware classification, terminal-offset checks, internal-anchor "
        "labels, and batch reporting provides a compact but detailed view of sample "
        "quality.",
    ),
    (
        "paragraph",
        "The current implementation has limitations that should be addressed before "
        "final submission. First, the example data demonstrate functionality but do "
        "not constitute a broad benchmark across protocols, chemistries, sequencers, "
        "or biological samples. Second, fixed-anchor approximate matching currently "
        "uses bounded Hamming distance after segment seeding; it does not model "
        "insertions and deletions as an edit-distance aligner would. Third, the "
        "best-orientation score is heuristic, and users should validate behavior "
        "for protocols with unusual or repeated motif structures. Fourth, the "
        "manuscript should include software tests, runtime benchmarks, public "
        "installation instructions, license information, and an archival DOI before "
        "submission.",
    ),
    (
        "paragraph",
        "Future development can extend scfastq-qc in several directions. Protocol "
        "templates could reduce configuration burden for common single-cell "
        "long-read assays. Edit-distance matching or minimizer-based motif search "
        "could improve indel tolerance. Barcode-aware modules could connect "
        "structure-level QC to cell-level summaries. Finally, integrating report "
        "outputs with workflow systems would make structure-aware QC easier to run "
        "as a routine step in reproducible sequencing pipelines.",
    ),
    ("heading1", "Availability and implementation"),
    (
        "paragraph",
        "The software is implemented in Python 3.10+ with an optional Rust accelerator. "
        "It is currently available in the accompanying source repository. Before "
        "submission, add the public GitHub URL, permanent archive DOI, software "
        "license, package version, installation channels, documentation URL, and "
        "example command lines verified from a clean environment.",
    ),
    ("heading1", "Data availability"),
    (
        "paragraph",
        "The demonstration FASTQ files and configuration used in this draft are "
        "included in the local repository under examples2/. If these data are not "
        "intended for public release, replace them with a public dataset or deposit "
        "an appropriate demonstration dataset in a stable repository before "
        "submission.",
    ),
    ("heading1", "Author contributions"),
    (
        "paragraph",
        "[To be completed according to the final author list. Suggested CRediT roles: "
        "Conceptualization, Methodology, Software, Validation, Visualization, "
        "Writing - original draft, Writing - review and editing.]",
    ),
    ("heading1", "Funding"),
    ("paragraph", "[To be completed.]"),
    ("heading1", "Conflict of interest"),
    ("paragraph", "The authors declare no competing interests. [Confirm before submission.]"),
    ("heading1", "References"),
    (
        "paragraph",
        "Andrews, S. (2010) FastQC: a quality control tool for high throughput "
        "sequence data. Babraham Bioinformatics. [Verify citation format and URL.]",
    ),
    (
        "paragraph",
        "De Coster, W. et al. (2018) NanoPack: visualizing and processing long-read "
        "sequencing data. Bioinformatics. [Verify final bibliographic details.]",
    ),
    (
        "paragraph",
        "Gao, Y., He, D. and Patro, R. (2026) QCatch: a framework for quality control "
        "assessment and analysis of single-cell sequencing data. Bioinformatics, "
        "42(5), btag184.",
    ),
    (
        "paragraph",
        "Oxford Nanopore Technologies. Pychopper: identify, orient and trim full-length "
        "cDNA reads. [Add versioned citation or repository archive if used as a "
        "direct comparator.]",
    ),
    (
        "paragraph",
        "[Add protocol-specific and single-cell long-read sequencing references "
        "matching the final biological use case.]",
    ),
    ("heading1", "Author revision checklist"),
    (
        "paragraph",
        "Before journal submission, replace placeholders, add a public software URL "
        "and license, provide versioned release metadata, run a clean-install test, "
        "add representative public datasets, benchmark runtime with and without the "
        "Rust accelerator, compare against relevant generic QC or protocol-specific "
        "preprocessing tools, and prepare final figures from the HTML report.",
    ),
]


def set_normal_style(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)
    for style_name in ["Heading 1", "Heading 2"]:
        styles[style_name].font.name = "Arial"
        styles[style_name].font.bold = True
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)


def add_paragraph(document: Document, text: str) -> None:
    for block in text.split("\n"):
        paragraph = document.add_paragraph(block)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.08


def add_table(document: Document, title: str, rows: list[list[str]]) -> None:
    document.add_paragraph(title).runs[0].bold = True
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = value
            if row_idx == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    document.add_paragraph()


def write_docx() -> None:
    OUTDIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    set_normal_style(document)
    for kind, *payload in MANUSCRIPT:
        if kind == "title":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(payload[0])
            run.bold = True
            run.font.size = Pt(16)
        elif kind == "heading1":
            document.add_heading(payload[0], level=1)
        elif kind == "heading2":
            document.add_heading(payload[0], level=2)
        elif kind == "paragraph":
            add_paragraph(document, payload[0])
        elif kind == "table":
            add_table(document, payload[0], payload[1])
    document.save(DOCX_PATH)


def write_markdown() -> None:
    lines: list[str] = []
    for kind, *payload in MANUSCRIPT:
        if kind == "title":
            lines.append(f"# {payload[0]}\n")
        elif kind == "heading1":
            lines.append(f"## {payload[0]}\n")
        elif kind == "heading2":
            lines.append(f"### {payload[0]}\n")
        elif kind == "paragraph":
            lines.append(payload[0] + "\n")
        elif kind == "table":
            title, rows = payload
            lines.append(f"**{title}**\n")
            header = rows[0]
            lines.append("| " + " | ".join(header) + " |")
            lines.append("| " + " | ".join("---" for _ in header) + " |")
            for row in rows[1:]:
                lines.append("| " + " | ".join(row) + " |")
            lines.append("")
    MD_PATH.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    write_docx()
    write_markdown()
    print(DOCX_PATH)
    print(MD_PATH)
